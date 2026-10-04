"""Genera la presentación final del Capstone (SpaceX Falcon 9) con python-pptx.

Uso:
    python presentation/generate_presentation.py          # solo el .pptx
    python presentation/generate_presentation.py --pdf    # .pptx + PDF para la entrega (requiere LibreOffice)

Lee las salidas de los notebooks y del script de capturas:
    * images/*.png        -> gráficas del EDA, mapas, dashboard y matrices de confusión
    * results/*.json      -> cifras (SQL, EDA, modelos, dashboard, distancias)

Si una salida aún no existe, la diapositiva lo indica con un recuadro "pendiente" en lugar de inventar
cifras. Las únicas excepciones son las consultas SQL y las métricas de los modelos: si todavía no se han
ejecutado los notebooks 04 y 07, se muestran los valores de referencia del laboratorio oficial de IBM y la
diapositiva lo declara en su nota de fuente.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------------------
# Configuración: cambia estos datos antes de entregar
# --------------------------------------------------------------------------------------
AUTHOR = "Diego Laureano"
REPO_URL = "https://github.com/alessandroxlc/spacex-falcon9-capstone"
DECK_DATE = date.today().strftime("%d/%m/%Y")

ROOT = Path(__file__).resolve().parent.parent
IMAGES = ROOT / "images"
RESULTS = ROOT / "results"
OUTPUT = ROOT / "presentation" / "SpaceX_Falcon9_Capstone.pptx"
PDF_NAME = "Data Science Capstone Project Report.pdf"

# --------------------------------------------------------------------------------------
# Identidad visual: azul noche (espacio) + naranja (llama del motor). Verde/rojo = aterrizó / no.
# --------------------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x12, 0x20)
NAVY_2 = RGBColor(0x16, 0x21, 0x36)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
SOFT = RGBColor(0xCB, 0xD5, 0xE1)
TINT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xF2, 0x6B, 0x1D)
GREEN = RGBColor(0x1F, 0xA6, 0x7A)
RED = RGBColor(0xD6, 0x45, 0x45)
FONT = "Calibri"
MONO = "Courier New"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_TOP = Inches(1.45)
CONTENT_W = W - 2 * MARGIN

# --------------------------------------------------------------------------------------
# Valores de referencia del laboratorio oficial de IBM (solo se usan si faltan los JSON)
# --------------------------------------------------------------------------------------
REFERENCE_SQL = [
    {"task": "Sitios de lanzamiento únicos", "columns": ["Launch_Site"],
     "rows": [["CCAFS LC-40"], ["VAFB SLC-4E"], ["KSC LC-39A"], ["CCAFS SLC-40"]]},
    {"task": "Sitios que empiezan por 'CCA'",
     "columns": ["Date", "Booster_Version", "Launch_Site", "Payload", "PAYLOAD_MASS__KG_"],
     "rows": [["2010-06-04", "F9 v1.0 B0003", "CCAFS LC-40", "Dragon Spacecraft Qualification Unit", "0"],
              ["2010-12-08", "F9 v1.0 B0004", "CCAFS LC-40", "Dragon demo flight C1", "0"],
              ["2012-05-22", "F9 v1.0 B0005", "CCAFS LC-40", "Dragon demo flight C2", "525"],
              ["2012-10-08", "F9 v1.0 B0006", "CCAFS LC-40", "SpaceX CRS-1", "500"],
              ["2013-03-01", "F9 v1.0 B0007", "CCAFS LC-40", "SpaceX CRS-2", "677"]]},
    {"task": "Masa total de carga NASA (CRS)", "columns": ["Total_Payload_Mass_kg"], "rows": [["45596"]]},
    {"task": "Masa media de carga F9 v1.1", "columns": ["Avg_Payload_Mass_kg"], "rows": [["2928.4"]]},
    {"task": "Primer aterrizaje exitoso en tierra", "columns": ["First_Successful_Ground_Landing"],
     "rows": [["2015-12-22"]]},
    {"task": "Éxito en barcaza con carga 4.000–6.000 kg", "columns": ["Booster_Version"],
     "rows": [["F9 FT B1022"], ["F9 FT B1026"], ["F9 FT B1021.2"], ["F9 FT B1031.2"]]},
    {"task": "Resultados de misión", "columns": ["Mission_Outcome", "Total"],
     "rows": [["Failure (in flight)", "1"], ["Success", "99"], ["Success (payload status unclear)", "1"]]},
    {"task": "Boosters con la carga máxima", "columns": ["Booster_Version", "PAYLOAD_MASS__KG_"],
     "rows": [[b, "15600"] for b in ["F9 B5 B1048.4", "F9 B5 B1049.4", "F9 B5 B1051.3", "F9 B5 B1056.4",
                                     "F9 B5 B1048.5", "F9 B5 B1051.4", "F9 B5 B1049.5", "F9 B5 B1060.2",
                                     "F9 B5 B1058.3", "F9 B5 B1051.6", "F9 B5 B1060.3", "F9 B5 B1049.7"]]},
    {"task": "Fallos en barcaza en 2015",
     "columns": ["Month", "Date", "Booster_Version", "Launch_Site", "Landing_Outcome"],
     "rows": [["01", "2015-01-10", "F9 v1.1 B1012", "CCAFS LC-40", "Failure (drone ship)"],
              ["04", "2015-04-14", "F9 v1.1 B1015", "CCAFS LC-40", "Failure (drone ship)"]]},
    {"task": "Ranking de aterrizajes 2010-06-04 a 2017-03-20", "columns": ["Landing_Outcome", "Total"],
     "rows": [["No attempt", "10"], ["Success (drone ship)", "5"], ["Failure (drone ship)", "5"],
              ["Success (ground pad)", "3"], ["Controlled (ocean)", "3"], ["Uncontrolled (ocean)", "2"],
              ["Failure (parachute)", "2"], ["Precluded (drone ship)", "1"]]},
]
SQL_QUERIES = [
    "SELECT DISTINCT Launch_Site\nFROM SPACEXTABLE;",
    "SELECT Date, Booster_Version, Launch_Site,\n       Payload, PAYLOAD_MASS__KG_\nFROM SPACEXTABLE\n"
    "WHERE Launch_Site LIKE 'CCA%'\nLIMIT 5;",
    "SELECT SUM(PAYLOAD_MASS__KG_)\nFROM SPACEXTABLE\nWHERE Customer = 'NASA (CRS)';",
    "SELECT AVG(PAYLOAD_MASS__KG_)\nFROM SPACEXTABLE\nWHERE Booster_Version = 'F9 v1.1';",
    "SELECT MIN(Date)\nFROM SPACEXTABLE\nWHERE Landing_Outcome =\n      'Success (ground pad)';",
    "SELECT Booster_Version\nFROM SPACEXTABLE\nWHERE Landing_Outcome =\n      'Success (drone ship)'\n"
    "  AND PAYLOAD_MASS__KG_ > 4000\n  AND PAYLOAD_MASS__KG_ < 6000;",
    "SELECT TRIM(Mission_Outcome),\n       COUNT(*) AS Total\nFROM SPACEXTABLE\nGROUP BY TRIM(Mission_Outcome);",
    "SELECT DISTINCT Booster_Version\nFROM SPACEXTABLE\nWHERE PAYLOAD_MASS__KG_ = (\n"
    "  SELECT MAX(PAYLOAD_MASS__KG_)\n  FROM SPACEXTABLE);",
    "SELECT substr(Date,6,2) AS Month,\n       Date, Booster_Version,\n       Launch_Site, Landing_Outcome\n"
    "FROM SPACEXTABLE\nWHERE Landing_Outcome =\n      'Failure (drone ship)'\n  AND substr(Date,1,4) = '2015';",
    "SELECT Landing_Outcome,\n       COUNT(*) AS Total\nFROM SPACEXTABLE\n"
    "WHERE Date BETWEEN '2010-06-04'\n               AND '2017-03-20'\nGROUP BY Landing_Outcome\nORDER BY Total DESC;",
]
SQL_TASKS_EN = [
    "All Launch Site Names", "Launch Site Names Begin with 'CCA'", "Total Payload Mass",
    "Average Payload Mass by F9 v1.1", "First Successful Ground Landing Date",
    "Successful Drone Ship Landing with Payload between 4000 and 6000",
    "Total Number of Successful and Failure Mission Outcomes", "Boosters Carried Maximum Payload",
    "2015 Launch Records", "Rank Landing Outcomes Between 2010-06-04 and 2017-03-20",
]
REFERENCE_MODELS = {
    "test_size": 18,
    "train_size": 72,
    "best_model": "Árbol de decisión",
    "models": {
        "Regresión logística": {"cv_accuracy": 0.8464, "test_accuracy": 0.8333,
                                "best_params": {"C": 0.01, "penalty": "l2", "solver": "lbfgs"},
                                "confusion_matrix": [[3, 3], [0, 12]]},
        "SVM": {"cv_accuracy": 0.8482, "test_accuracy": 0.8333,
                "best_params": {"C": 1.0, "gamma": 0.0316, "kernel": "sigmoid"},
                "confusion_matrix": [[3, 3], [0, 12]]},
        "Árbol de decisión": {"cv_accuracy": 0.8750, "test_accuracy": 0.8333,
                              "best_params": {"criterion": "gini", "max_depth": 4},
                              "confusion_matrix": [[3, 3], [0, 12]]},
        "KNN": {"cv_accuracy": 0.8482, "test_accuracy": 0.8333,
                "best_params": {"algorithm": "auto", "n_neighbors": 10, "p": 1},
                "confusion_matrix": [[3, 3], [0, 12]]},
    },
}
NB = {  # enlaces a cada notebook del repositorio
    "api": "notebooks/01_data_collection_api.ipynb",
    "scraping": "notebooks/02_data_collection_webscraping.ipynb",
    "wrangling": "notebooks/03_data_wrangling.ipynb",
    "sql": "notebooks/04_eda_sql.ipynb",
    "viz": "notebooks/05_eda_visualization.ipynb",
    "folium": "notebooks/06_interactive_map_folium.ipynb",
    "ml": "notebooks/07_machine_learning_prediction.ipynb",
    "dash": "dashboard/spacex_dash_app.py",
}


def load_json(name: str):
    path = RESULTS / name
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def link(key: str) -> str:
    return f"{REPO_URL}/blob/main/{NB[key]}"


def pct(value: float | None, digits: int = 1) -> str:
    return "—" if value is None else f"{value * 100:.{digits}f}".replace(".", ",") + " %"


# ======================================================================================
# Primitivas de dibujo
# ======================================================================================
class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    # ---------------------------------------------------------------- fondos y marcos
    def new_slide(self, dark: bool = False, notes: str = ""):
        slide = self.prs.slides.add_slide(self.blank)
        self.count += 1
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY if dark else WHITE
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def content_slide(self, title: str, notes: str = "", source: str | None = None, kicker: str | None = None):
        """Diapositiva de contenido. `kicker` repite el nombre de la diapositiva en la plantilla oficial de
        IBM (en inglés) para que el evaluador encuentre cada sección del rúbrica sin ambigüedad."""
        slide = self.new_slide(notes=notes)
        if kicker:
            self.text(slide, kicker.upper(), MARGIN, Inches(0.3), CONTENT_W, Inches(0.3), size=12, bold=True,
                      color=ORANGE)
        self.text(slide, title, MARGIN, Inches(0.55), CONTENT_W, Inches(0.75), size=30, bold=True, color=NAVY,
                  anchor=MSO_ANCHOR.MIDDLE)
        self.footer(slide, source)
        return slide

    def footer(self, slide, source: str | None = None, dark: bool = False) -> None:
        color = SOFT if dark else MUTED
        if source:
            self.text(slide, source, MARGIN, Inches(6.92), Inches(7.6), Inches(0.35), size=10, color=color,
                      italic=True)
        self.text(slide, f"{AUTHOR} · SpaceX Falcon 9 · IBM Data Science Capstone   {self.count}", Inches(8.4),
                  Inches(6.92), Inches(4.33), Inches(0.35), size=10, color=color, align=PP_ALIGN.RIGHT)

    # ---------------------------------------------------------------- texto
    def text(self, slide, content, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, font=FONT, url: str | None = None):
        """Cuadro de texto. `content` puede ser str o lista de str (un párrafo por elemento)."""
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = content if isinstance(content, list) else [content]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = line
            f = run.font
            f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
            f.color.rgb = color
            if url:
                run.hyperlink.address = url
        return box

    def rich(self, slide, x, y, w, h, paragraphs, size=16, color=INK, space_after=8, anchor=MSO_ANCHOR.TOP):
        """Párrafos con viñetas y tramos en negrita.

        Cada párrafo es una lista de tramos (texto, negrita) o un dict {"runs": [...], "bullet": bool}.
        """
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, para in enumerate(paragraphs):
            if isinstance(para, str):
                para = {"runs": [(para, False)], "bullet": True}
            elif isinstance(para, list):
                para = {"runs": para, "bullet": True}
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space_after)
            for txt, is_bold in para["runs"]:
                run = p.add_run()
                run.text = txt
                run.font.name, run.font.size, run.font.bold = FONT, Pt(para.get("size", size)), is_bold
                run.font.color.rgb = para.get("color", color)
            if para.get("bullet", True):
                _bullet(p)
        return box

    # ---------------------------------------------------------------- formas
    def box(self, slide, x, y, w, h, fill=TINT, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(1.25)
        shp.shadow.inherit = False
        shp.text_frame.text = ""
        return shp

    def badge(self, slide, label: str, x, y, d=Inches(0.62), fill=ORANGE, color=WHITE, size=20):
        """Círculo numerado: el motivo visual de la presentación (como un parche de misión)."""
        circ = self.box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.name, run.font.size, run.font.bold = FONT, Pt(size), True
        run.font.color.rgb = color
        return circ

    def arrow(self, slide, x1, y1, x2, y2, color=SOFT):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        conn.line.color.rgb = color
        conn.line.width = Pt(2)
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        return conn

    def stat(self, slide, value: str, label: str, x, y, w, color=ORANGE, label_color=INK, size=44):
        self.text(slide, value, x, y, w, Inches(0.9), size=size, bold=True, color=color)
        self.text(slide, label, x, y + Inches(0.9), w, Inches(0.7), size=14, color=label_color)

    def link_chip(self, slide, key: str, x, y, w=CONTENT_W):
        self.text(slide, f"GitHub: {link(key)}", x, y, w, Inches(0.3), size=12, color=ORANGE, url=link(key))

    # ---------------------------------------------------------------- imágenes
    def image(self, slide, filename: str, x, y, w, h, hint: str):
        """Inserta images/<filename> ajustada a la caja; si no existe dibuja un recuadro 'pendiente'."""
        path = IMAGES / filename
        if not path.exists():
            frame = self.box(slide, x, y, w, h, fill=TINT, line=SOFT)
            frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            self.text(slide, ["Figura pendiente", hint], x + Inches(0.4), y, w - Inches(0.8), h, size=14,
                      color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            return None
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)

    # ---------------------------------------------------------------- tablas
    def table(self, slide, columns, rows, x, y, w, col_widths=None, size=12, row_h=Inches(0.36)):
        n_rows = len(rows) + 1
        shape = slide.shapes.add_table(n_rows, len(columns), x, y, w, row_h * n_rows)
        tbl = shape.table
        tbl.first_row = True
        # Quita el estilo por defecto (bandas azules de Office)
        tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
        style = tbl_pr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
        if col_widths:
            total = sum(col_widths)
            for i, cw in enumerate(col_widths):
                tbl.columns[i].width = int(w * cw / total)
        for r in range(n_rows):
            tbl.rows[r].height = row_h
            values = columns if r == 0 else rows[r - 1]
            for c, value in enumerate(values):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if r == 0 else (TINT if r % 2 == 0 else WHITE)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = str(value)
                run.font.name, run.font.size = FONT, Pt(size)
                run.font.bold = r == 0
                run.font.color.rgb = WHITE if r == 0 else INK
                _cell_border(cell, "CBD5E1")
        return shape

    def code(self, slide, sql: str, x, y, w, h):
        self.box(slide, x, y, w, h, fill=NAVY_2, radius=0.04)
        self.text(slide, sql.split("\n"), x + Inches(0.3), y + Inches(0.28), w - Inches(0.6), h - Inches(0.5),
                  size=14, color=WHITE, font=MONO)

    # ---------------------------------------------------------------- gráficos nativos
    def bar_chart(self, slide, categories, series: dict, x, y, w, h, percent=True, colors=None, legend=False,
                  y_max=1.0):
        data = CategoryChartData()
        data.categories = categories
        for name, values in series.items():
            data.add_series(name, values)
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data).chart
        _style_chart(chart, legend)
        va = chart.value_axis
        va.maximum_scale, va.minimum_scale = y_max, 0
        va.tick_labels.number_format = "0%" if percent else "General"
        va.tick_labels.number_format_is_linked = False
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = "0%" if percent else "General"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size, dl.font.color.rgb, dl.font.name = Pt(12), INK, FONT
        for s, color in zip(plot.series, colors or [ORANGE, MUTED]):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = color
        return chart

    def line_chart(self, slide, categories, values, x, y, w, h, name="Tasa de éxito"):
        data = CategoryChartData()
        data.categories = categories
        data.add_series(name, values)
        chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, data).chart
        _style_chart(chart, legend=False)
        va = chart.value_axis
        va.maximum_scale, va.minimum_scale = 1.0, 0
        va.tick_labels.number_format = "0%"
        va.tick_labels.number_format_is_linked = False
        s = chart.plots[0].series[0]
        s.format.line.color.rgb = ORANGE
        s.format.line.width = Pt(3)
        s.smooth = False
        s.marker.size = 9
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = ORANGE
        s.marker.format.line.color.rgb = ORANGE
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.position = XL_LABEL_POSITION.ABOVE
        dl.font.size, dl.font.color.rgb, dl.font.name = Pt(12), INK, FONT
        return chart

    def save(self, path: Path) -> None:
        path.parent.mkdir(exist_ok=True)
        theme = self.prs.slide_masters[0].part.part_related_by(RT.THEME)
        xml = etree.fromstring(theme.blob)
        for tag in ("a:hlink", "a:folHlink"):
            node = xml.find(f".//{qn(tag)}")
            for child in list(node):
                node.remove(child)
            etree.SubElement(node, qn("a:srgbClr")).set("val", "F26B1D")
        theme._blob = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
        self.prs.save(path)


def _bullet(paragraph, color="F26B1D") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.3))))
    pPr.set("indent", str(-Emu(Inches(0.3))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def _cell_border(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tc_pr, qn(tag), w="9525", cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex_color)
    # tcPr exige que los bordes vayan antes del relleno
    fill = tc_pr.find(qn("a:solidFill"))
    if fill is not None:
        tc_pr.remove(fill)
        tc_pr.append(fill)


def _style_chart(chart, legend: bool) -> None:
    chart.font.name = FONT
    chart.font.size = Pt(12)
    chart.font.color.rgb = MUTED
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    va.format.line.fill.background()
    chart.category_axis.format.line.color.rgb = SOFT
    chart.category_axis.has_major_gridlines = False


# ======================================================================================
# Contenido
# ======================================================================================
def build() -> Path:
    sql = load_json("sql_results.json")
    sql_source = "Fuente: results/sql_results.json (notebook 04)"
    if not sql:
        sql = REFERENCE_SQL
        sql_source = "Fuente: valores de referencia del laboratorio IBM (ejecuta el notebook 04 para regenerarlos)"
    models = load_json("model_results.json")
    model_source = "Fuente: results/model_results.json (notebook 07)"
    if not models:
        models = REFERENCE_MODELS
        model_source = "Fuente: valores de referencia del laboratorio IBM (ejecuta el notebook 07 para regenerarlos)"
    eda = load_json("eda.json")
    wrangling = load_json("data_wrangling.json")
    folium_res = load_json("folium.json")
    dash = load_json("dashboard.json")

    best = models["best_model"]
    best_m = models["models"][best]
    test_scores = [m["test_accuracy"] for m in models["models"].values()]
    tie = test_scores.count(max(test_scores)) > 1
    criterion = "desempate por validación cruzada" if tie else "mayor exactitud en prueba"
    cv_scores = [m["cv_accuracy"] for m in models["models"].values()]
    cv_range = f"{pct(min(cv_scores), 0)}–{pct(max(cv_scores), 0)}"
    point = f"{100 / models['test_size']:.1f}".replace(".", ",")
    test_acc = best_m["test_accuracy"]
    n_launches = (wrangling or eda or {}).get("rows")
    success_rate = (wrangling or {}).get("success_rate")

    d = Deck()

    # ------------------------------------------------------------------ 1. Portada
    s = d.new_slide(dark=True, notes=(
        "Presentación del proyecto final. En una frase: usamos datos públicos de lanzamientos de SpaceX para "
        "predecir si la primera etapa del cohete Falcon 9 aterrizará, porque de eso depende el precio de un "
        "lanzamiento."))
    d.badge(s, "F9", MARGIN, Inches(1.3), d=Inches(1.0), size=26)
    d.text(s, "IBM APPLIED DATA SCIENCE CAPSTONE", MARGIN, Inches(2.65), Inches(10), Inches(0.4), size=14,
           bold=True, color=ORANGE)
    d.text(s, "¿Aterrizará la primera etapa?", MARGIN, Inches(3.05), Inches(11.5), Inches(1.0), size=48,
           bold=True, color=WHITE)
    d.text(s, "Predicción del aterrizaje del SpaceX Falcon 9 con ciencia de datos y machine learning",
           MARGIN, Inches(4.1), Inches(11), Inches(0.6), size=22, color=SOFT)
    d.text(s, [AUTHOR, DECK_DATE], MARGIN, Inches(5.6), Inches(6), Inches(0.8), size=16, color=WHITE)
    d.text(s, f"Repositorio GitHub: {REPO_URL}", MARGIN, Inches(6.5), Inches(10), Inches(0.4), size=14,
           color=ORANGE, url=REPO_URL)

    # ------------------------------------------------------------------ 2. Índice
    s = d.content_slide("Índice", kicker="Outline", notes="Recorrido de la presentación siguiendo la estructura del curso.")
    items = ["Resumen ejecutivo", "Introducción", "Metodología", "Resultados",
             "Hallazgos del EDA (visualización y SQL)", "Análisis de proximidad con Folium",
             "Dashboard con Plotly Dash", "Análisis predictivo (clasificación)", "Conclusiones", "Apéndice"]
    for i, item in enumerate(items):
        col, row = divmod(i, 5)
        x = MARGIN + col * Inches(6.1)
        y = CONTENT_TOP + Inches(0.25) + row * Inches(1.0)
        d.badge(s, str(i + 1), x, y, d=Inches(0.55), size=16)
        d.text(s, item, x + Inches(0.8), y, Inches(5.0), Inches(0.55), size=20, color=INK,
               anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------------------------ 3. Resumen ejecutivo
    s = d.content_slide("Resumen ejecutivo", kicker="Executive Summary", notes=(
        "Mensaje clave: con datos públicos y modelos sencillos se puede anticipar si el booster aterrizará en "
        "aproximadamente 8 de cada 10 lanzamientos de prueba. Explicar brevemente las fases de la metodología."),
        source=model_source)
    d.text(s, "Metodología", MARGIN, CONTENT_TOP, Inches(7.2), Inches(0.45), size=20, bold=True, color=NAVY)
    d.rich(s, MARGIN, CONTENT_TOP + Inches(0.6), Inches(7.2), Inches(4.6), [
        [("Recolección: ", True), ("API REST de SpaceX y web scraping de Wikipedia.", False)],
        [("Preparación: ", True), ("limpieza, imputación y etiqueta binaria Class (1 = aterrizó).", False)],
        [("EDA: ", True), ("gráficos con Seaborn y 10 consultas SQL.", False)],
        [("Analítica interactiva: ", True), ("mapas con Folium y dashboard con Plotly Dash.", False)],
        [("Modelado: ", True), ("regresión logística, SVM, árbol de decisión y KNN con GridSearchCV.", False)],
    ], size=16, space_after=12)
    panel_x = Inches(8.3)
    d.box(s, panel_x, CONTENT_TOP, Inches(4.43), Inches(5.1), fill=NAVY)
    d.stat(s, pct(test_acc, 2), f"exactitud en prueba del mejor modelo ({best.lower()})", panel_x + Inches(0.4),
           CONTENT_TOP + Inches(0.3), Inches(3.8), label_color=WHITE)
    d.stat(s, "4", "modelos de clasificación comparados", panel_x + Inches(0.4), CONTENT_TOP + Inches(1.9),
           Inches(3.8), color=WHITE, label_color=SOFT)
    d.stat(s, str(n_launches) if n_launches else "—", "lanzamientos Falcon 9 analizados (2010–2020)",
           panel_x + Inches(0.4), CONTENT_TOP + Inches(3.5), Inches(3.8), color=WHITE, label_color=SOFT)

    # ------------------------------------------------------------------ 4. Introducción
    s = d.content_slide("Introducción: el valor de reutilizar el cohete", kicker="Introduction", notes=(
        "SpaceX anuncia lanzamientos del Falcon 9 por unos 62 millones de dólares; otros proveedores superan "
        "los 165 millones. La diferencia se explica porque SpaceX reutiliza la primera etapa. Si sabemos "
        "predecir si aterrizará, podemos estimar el coste de un lanzamiento: útil para una empresa que quiera "
        "competir con SpaceX en una licitación."), source="Fuente de precios: enunciado del curso IBM Applied Data Science Capstone")
    d.box(s, MARGIN, CONTENT_TOP, Inches(3.6), Inches(2.3), fill=TINT)
    d.stat(s, "62 M$", "precio publicado por SpaceX para un Falcon 9", MARGIN + Inches(0.35),
           CONTENT_TOP + Inches(0.3), Inches(3.0), color=GREEN)
    d.box(s, MARGIN + Inches(3.9), CONTENT_TOP, Inches(3.6), Inches(2.3), fill=TINT)
    d.stat(s, "165 M$+", "coste habitual de otros proveedores", MARGIN + Inches(4.25),
           CONTENT_TOP + Inches(0.3), Inches(3.0), color=RED)
    d.text(s, "La diferencia: SpaceX recupera y reutiliza la primera etapa del cohete.", MARGIN,
           CONTENT_TOP + Inches(2.6), Inches(7.5), Inches(0.5), size=18, bold=True, color=NAVY)
    d.rich(s, MARGIN, CONTENT_TOP + Inches(3.3), Inches(7.5), Inches(2.0), [
        "Si el booster aterriza, el lanzamiento es más barato.",
        "Predecir el aterrizaje permite estimar el coste de cada misión.",
        "Una empresa competidora (como la ficticia SpaceY) puede usarlo para preparar sus ofertas.",
    ], size=16)
    d.text(s, "Preguntas a responder", Inches(8.6), CONTENT_TOP, Inches(4.1), Inches(0.45), size=20, bold=True,
           color=NAVY)
    for i, q in enumerate([
        "¿Qué variables influyen en el éxito del aterrizaje?",
        "¿Cómo ha evolucionado la tasa de éxito con los años?",
        "¿Influye la ubicación del sitio de lanzamiento?",
        "¿Con qué exactitud podemos predecir el resultado?",
    ]):
        y = CONTENT_TOP + Inches(0.7) + i * Inches(1.1)
        d.badge(s, str(i + 1), Inches(8.6), y, d=Inches(0.5), size=14)
        d.text(s, q, Inches(9.3), y - Inches(0.05), Inches(3.4), Inches(0.9), size=15, color=INK)

    # ------------------------------------------------------------------ Sección 1: Metodología
    section(d, "1", "Metodología", "Cómo se obtuvieron, prepararon y analizaron los datos")

    s = d.content_slide("Resumen de la metodología", kicker="Methodology · Executive Summary", notes=(
        "El proyecto es una tubería de datos completa: obtener datos, limpiarlos, explorarlos, visualizarlos de "
        "forma interactiva y finalmente entrenar modelos que predicen el aterrizaje."))
    steps = [
        ("Recolección", "API REST de SpaceX y web scraping de Wikipedia"),
        ("Wrangling", "Limpieza, imputación de nulos y etiqueta Class"),
        ("EDA", "Visualización (Seaborn) y consultas SQL"),
        ("Mapas", "Análisis geoespacial con Folium"),
        ("Dashboard", "Aplicación interactiva con Plotly Dash"),
        ("Modelado", "4 clasificadores + GridSearchCV + evaluación"),
    ]
    card_w = Inches(1.85)
    gap = (CONTENT_W - 6 * card_w) / 5
    for i, (name, desc) in enumerate(steps):
        x = MARGIN + i * (card_w + gap)
        y = CONTENT_TOP + Inches(0.8)
        d.box(s, x, y, card_w, Inches(3.4), fill=TINT)
        d.badge(s, str(i + 1), x + (card_w - Inches(0.7)) // 2, y + Inches(0.35), d=Inches(0.7))
        d.text(s, name, x + Inches(0.1), y + Inches(1.3), card_w - Inches(0.2), Inches(0.5), size=16,
               bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        d.text(s, desc, x + Inches(0.15), y + Inches(1.9), card_w - Inches(0.3), Inches(1.4), size=14,
               color=INK, align=PP_ALIGN.CENTER)
        if i < 5:
            d.arrow(s, x + card_w + Inches(0.03), y + Inches(0.7), x + card_w + gap - Inches(0.03), y + Inches(0.7))
    d.text(s, "Herramientas: Python · pandas · NumPy · requests · BeautifulSoup · SQLite · Matplotlib · Seaborn · "
              "Folium · Plotly Dash · scikit-learn", MARGIN, Inches(6.05), CONTENT_W, Inches(0.5), size=14,
           color=MUTED, align=PP_ALIGN.CENTER)

    s = d.content_slide("Recolección de datos", kicker="Methodology · Data Collection", notes=(
        "Dos fuentes complementarias: la API de SpaceX (datos estructurados en JSON) y la tabla de lanzamientos "
        "de Wikipedia (HTML). Ambas se convierten en tablas de pandas."))
    for i, (title, lines) in enumerate([
        ("SpaceX REST API v4", ["Endpoints: launches, rockets, launchpads, payloads y cores.",
                                "Respuesta en JSON → pd.json_normalize.",
                                "Resultado: data/dataset_part_1.csv"]),
        ("Web scraping (Wikipedia)", ["Página “List of Falcon 9 and Falcon Heavy launches”.",
                                      "requests + BeautifulSoup sobre las tablas HTML.",
                                      "Resultado: data/spacex_web_scraped.csv"]),
    ]):
        x = MARGIN + i * Inches(6.2)
        d.box(s, x, CONTENT_TOP + Inches(0.2), Inches(5.9), Inches(3.4), fill=TINT)
        d.badge(s, "A" if i == 0 else "B", x + Inches(0.35), CONTENT_TOP + Inches(0.5))
        d.text(s, title, x + Inches(1.2), CONTENT_TOP + Inches(0.5), Inches(4.4), Inches(0.62), size=22,
               bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        d.rich(s, x + Inches(0.4), CONTENT_TOP + Inches(1.45), Inches(5.2), Inches(2.0), lines, size=16)
    d.text(s, "Ambas fuentes se unifican en un dataset de un lanzamiento por fila, con variables de la misión "
              "(fecha, booster, carga, órbita, sitio) y el resultado del aterrizaje.", MARGIN, Inches(5.35),
           CONTENT_W, Inches(0.9), size=16, color=INK)

    flow_slide(d, "Recopilación de datos: API de SpaceX", [
        "GET /launches/past\n(copia estática del curso)", "pd.json_normalize()", "Filtrar 1 núcleo, 1 carga\ny fecha ≤ 13-11-2020",
        "Resolver IDs en /rockets,\n/launchpads, /payloads, /cores", "Filtrar Falcon 9 e\nimputar PayloadMass (media)",
        "dataset_part_1.csv",
    ], "api", kicker="Methodology · Data Collection – SpaceX API", notes=(
        "Se descarga el histórico de lanzamientos, se filtran las misiones simples y se completan los datos "
        "consultando los demás endpoints. Al final se queda solo el Falcon 9 y se rellenan las masas que faltan "
        "con la media."))

    flow_slide(d, "Recopilación de datos: web scraping", [
        "requests.get(URL de Wikipedia\noldid=1027686922)", "BeautifulSoup(html)", "find_all('table',\n'wikitable plainrowheaders')",
        "Extraer nombres de\ncolumna (<th>)", "Recorrer filas y limpiar\nceldas (fecha, booster, masa…)",
        "spacex_web_scraped.csv",
    ], "scraping", kicker="Methodology · Data Collection – Scraping", notes=(
        "Se descarga una versión fija de la página de Wikipedia para que el resultado sea reproducible, se buscan "
        "las tablas de lanzamientos y se extrae cada fila a un diccionario que luego se convierte en DataFrame."))

    s = d.content_slide("Metodología de limpieza de datos (data wrangling)", kicker="Methodology · Data Wrangling", notes=(
        "La columna Outcome combina si el aterrizaje tuvo éxito y dónde se intentó. La convertimos en una "
        "etiqueta binaria: 1 si aterrizó, 0 si no. Es la variable que el modelo aprenderá a predecir."),
        source=f"Tasa global de éxito tras el wrangling: {pct(success_rate)}" if success_rate is not None else None)
    d.rich(s, MARGIN, CONTENT_TOP, Inches(5.6), Inches(4.4), [
        [("Exploración: ", True), ("porcentaje de nulos, tipos de dato y recuentos por sitio y órbita.", False)],
        [("Nulos: ", True), ("PayloadMass se imputa con la media; LandingPad se conserva (nulo = sin plataforma).", False)],
        [("Etiqueta: ", True), ("Class = 1 si Outcome empieza por “True”; 0 en cualquier otro caso.", False)],
        [("Salida: ", True), ("data/dataset_part_2.csv", False)],
    ], size=16, space_after=14)
    d.link_chip(s, "wrangling", MARGIN, Inches(6.4))
    d.table(s, ["Outcome", "Significado", "Class"], [
        ["True ASDS", "Aterrizó en barcaza", "1"], ["True RTLS", "Aterrizó en tierra", "1"],
        ["True Ocean", "Descenso controlado al océano", "1"], ["False ASDS", "Falló en barcaza", "0"],
        ["False RTLS", "Falló en tierra", "0"], ["False Ocean", "Falló sobre el océano", "0"],
        ["None None / None ASDS", "Sin intento de aterrizaje", "0"],
    ], Inches(6.6), CONTENT_TOP, Inches(6.13), col_widths=[2.2, 3.2, 0.8], size=14, row_h=Inches(0.5))

    s = d.content_slide("EDA con visualización de datos", kicker="Methodology · EDA with Data Visualization", notes=(
        "Cada gráfico responde una pregunta concreta. Los diagramas de dispersión muestran relaciones entre dos "
        "variables, la barra compara tasas entre categorías y la línea muestra la evolución en el tiempo."))
    d.table(s, ["Gráfico", "Variables", "Pregunta que responde"], [
        ["Dispersión", "N.º de vuelo vs. sitio", "¿Mejora el éxito con la experiencia en cada sitio?"],
        ["Dispersión", "Carga útil vs. sitio", "¿Hay sitios para cargas pesadas?"],
        ["Barras", "Tasa de éxito por órbita", "¿Qué órbitas son más fáciles?"],
        ["Dispersión", "N.º de vuelo vs. órbita", "¿Cambia la relación con la experiencia según la órbita?"],
        ["Dispersión", "Carga útil vs. órbita", "¿Influye la masa según la órbita?"],
        ["Línea", "Tasa de éxito por año", "¿Cuál es la tendencia en el tiempo?"],
    ], MARGIN, CONTENT_TOP, CONTENT_W, col_widths=[1.4, 2.6, 4.6], size=16, row_h=Inches(0.58))
    d.text(s, "Después se aplica One-Hot Encoding a Orbit, LaunchSite, LandingPad y Serial → dataset_part_3.csv",
           MARGIN, Inches(5.75), CONTENT_W, Inches(0.4), size=16, color=INK)
    d.link_chip(s, "viz", MARGIN, Inches(6.4))

    s = d.content_slide("EDA con SQL", kicker="Methodology · EDA with SQL", notes=(
        "Se cargó el dataset en una base SQLite y se respondieron diez preguntas con SQL: filtros con WHERE y "
        "LIKE, agregaciones con SUM, AVG, MIN y COUNT, agrupaciones con GROUP BY y una subconsulta."))
    queries = ["Sitios de lanzamiento únicos (DISTINCT)", "5 registros de sitios que empiezan por 'CCA' (LIKE)",
               "Masa total para NASA (CRS) (SUM)", "Masa media del booster F9 v1.1 (AVG)",
               "Primer aterrizaje exitoso en tierra (MIN)", "Éxito en barcaza con carga 4.000–6.000 kg",
               "Misiones exitosas y fallidas (GROUP BY)", "Boosters con la carga máxima (subconsulta)",
               "Fallos en barcaza en 2015 (substr)", "Ranking de aterrizajes 2010–2017 (ORDER BY)"]
    for i, q in enumerate(queries):
        col, row = divmod(i, 5)
        x = MARGIN + col * Inches(6.15)
        y = CONTENT_TOP + Inches(0.1) + row * Inches(0.88)
        d.badge(s, str(i + 1), x, y, d=Inches(0.5), size=14)
        d.text(s, q, x + Inches(0.7), y, Inches(5.3), Inches(0.5), size=16, anchor=MSO_ANCHOR.MIDDLE)
    d.link_chip(s, "sql", MARGIN, Inches(6.4))

    s = d.content_slide("Análisis visual interactivo: mapa con Folium", kicker="Methodology · Interactive Visual Analytics (Folium)", notes=(
        "Folium crea mapas interactivos en HTML. Marcamos los sitios, agrupamos los lanzamientos con colores "
        "según el resultado y medimos distancias a la costa, el ferrocarril, la autopista y la ciudad más cercana."))
    for i, (obj, why) in enumerate([
        ("Círculos y etiquetas (folium.Circle, DivIcon)", "Ubicar los 4 sitios y ver su cercanía al ecuador y a la costa."),
        ("MarkerCluster verde / rojo", "Ver qué sitios acumulan más aterrizajes exitosos."),
        ("MousePosition", "Leer coordenadas de puntos cercanos directamente en el mapa."),
        ("PolyLine + distancia haversine", "Medir la distancia a la costa, al ferrocarril, a la autopista y a la ciudad."),
    ]):
        y = CONTENT_TOP + Inches(0.1) + i * Inches(1.1)
        d.badge(s, str(i + 1), MARGIN, y, d=Inches(0.55), size=16)
        d.text(s, obj, MARGIN + Inches(0.8), y - Inches(0.05), Inches(11), Inches(0.4), size=18, bold=True,
               color=NAVY)
        d.text(s, why, MARGIN + Inches(0.8), y + Inches(0.38), Inches(11), Inches(0.4), size=16)
    d.link_chip(s, "folium", MARGIN, Inches(6.05))
    d.link_chip(s, "dash", MARGIN, Inches(6.4))

    s = d.content_slide("Análisis visual interactivo: dashboard con Plotly Dash", kicker="Methodology · Interactive Visual Analytics (Plotly Dash)", notes=(
        "El dashboard permite explorar los datos sin programar: eliges un sitio y un rango de masa y los gráficos "
        "se actualizan solos mediante callbacks."))
    for i, (comp, why) in enumerate([
        ("Desplegable de sitios (dcc.Dropdown)", "Filtrar todos los sitios o uno concreto."),
        ("Gráfico de pastel", "Todos: cuota de éxitos por sitio. Un sitio: éxito vs. fallo."),
        ("Control de rango (dcc.RangeSlider)", "Elegir el rango de masa de carga útil (0–10.000 kg)."),
        ("Dispersión carga vs. resultado", "Ver la relación masa–éxito, coloreada por versión del booster."),
    ]):
        col, row = divmod(i, 2)
        x = MARGIN + col * Inches(6.15)
        y = CONTENT_TOP + Inches(0.1) + row * Inches(2.2)
        d.box(s, x, y, Inches(5.9), Inches(1.95), fill=TINT)
        d.badge(s, str(i + 1), x + Inches(0.3), y + Inches(0.3), d=Inches(0.55), size=16)
        d.text(s, comp, x + Inches(1.05), y + Inches(0.3), Inches(4.6), Inches(0.55), size=18, bold=True,
               color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, why, x + Inches(1.05), y + Inches(0.95), Inches(4.6), Inches(0.9), size=16)
    d.link_chip(s, "folium", MARGIN, Inches(6.05))
    d.link_chip(s, "dash", MARGIN, Inches(6.4))

    flow_slide(d, "Análisis predictivo (clasificación)", [
        "Cargar X (one-hot)\ny Y = Class", "StandardScaler", "train_test_split\n80/20, random_state=2",
        "GridSearchCV (cv=10)\nLogReg · SVM · Árbol · KNN", "Exactitud en prueba\ny matriz de confusión",
        "Elegir el mejor\nmodelo",
    ], "ml", kicker="Methodology · Predictive Analysis (Classification)", notes=(
        "Proceso estándar de aprendizaje supervisado: estandarizar, separar datos de prueba que el modelo nunca "
        "ve, buscar los mejores hiperparámetros con validación cruzada y comparar en el conjunto de prueba."))

    s = d.content_slide("Resultados", kicker="Results", notes=(
        "Vista previa de los tres bloques de resultados que se detallan a continuación."), source=model_source)
    blocks = [
        ("EDA", ["La tasa de éxito crece con el número de vuelo y con los años.",
                 "Algunas órbitas (SSO, HEO, GEO, ES-L1) tienen 100 % de éxito con pocos vuelos; GTO ronda el 50 %."]),
        ("Analítica interactiva", ["Todos los sitios están junto a la costa y cerca de infraestructura de transporte.",
                                   "El dashboard compara sitios y rangos de carga en segundos."]),
        ("Análisis predictivo", [f"Exactitud en prueba de {pct(test_acc, 2)} en {models['test_size']} lanzamientos.",
                                 f"Mejor modelo: {best} ({criterion}).",
                                 f"En validación cruzada los 4 modelos quedan entre {cv_range}."]),
    ]
    for i, (title, lines) in enumerate(blocks):
        x = MARGIN + i * Inches(4.1)
        d.box(s, x, CONTENT_TOP, Inches(3.83), Inches(4.9), fill=NAVY if i == 2 else TINT)
        d.text(s, title, x + Inches(0.35), CONTENT_TOP + Inches(0.35), Inches(3.2), Inches(0.5), size=20, bold=True,
               color=ORANGE if i == 2 else NAVY)
        d.rich(s, x + Inches(0.35), CONTENT_TOP + Inches(1.1), Inches(3.2), Inches(3.6), lines, size=16,
               color=WHITE if i == 2 else INK, space_after=12)

    # ------------------------------------------------------------------ Sección 2: EDA
    section(d, "2", "Hallazgos del EDA", "Visualización de datos y consultas SQL")

    viz_hint = "Se genera al ejecutar notebooks/05_eda_visualization.ipynb"
    for title, img, takeaway, notes, kicker in [
        ("Número de vuelo vs. sitio de lanzamiento", "eda_flight_vs_site.png",
         ["Los primeros vuelos, casi todos fallidos, salieron de CCAFS SLC-40.",
          "En todos los sitios, los vuelos más recientes aterrizan con más frecuencia.",
          "La experiencia acumulada mejora el resultado."],
         "Cada punto es un lanzamiento. Verde: aterrizó; rojo: no. Hacia la derecha (más experiencia) predomina el verde.",
         "EDA Results · Flight Number vs. Launch Site"),
        ("Carga útil vs. sitio de lanzamiento", "eda_payload_vs_site.png",
         ["VAFB SLC-4E no lanza cargas muy pesadas (> 10.000 kg).",
          "Las cargas pesadas salen de Florida y suelen aterrizar con éxito.",
          "Por debajo de ~7.000 kg se mezclan éxitos y fallos."],
         "Relación entre la masa de la carga y el sitio. No hay una relación simple: la masa por sí sola no decide.",
         "EDA Results · Payload vs. Launch Site"),
    ]:
        eda_image_slide(d, title, img, takeaway, viz_hint, notes, kicker)

    s = d.content_slide("Tasa de éxito por tipo de órbita", kicker="EDA Results · Success Rate vs. Orbit Type", notes=(
        "La barra muestra el porcentaje de aterrizajes exitosos en cada órbita. Ojo: algunas órbitas con 100 % "
        "tienen solo uno o dos lanzamientos, así que no son concluyentes."),
        source="Fuente: results/eda.json (notebook 05)" if eda else None)
    if eda:
        rates = eda["success_rate_by_orbit"]
        cats = list(rates)
        d.bar_chart(s, cats, {"Tasa de éxito": [rates[c] for c in cats]}, MARGIN, CONTENT_TOP, Inches(8.2),
                    Inches(5.2), y_max=1.15)
        counts = eda.get("launches_by_orbit", {})
        top = max(counts, key=counts.get) if counts else None
        side = ["ES-L1, GEO, HEO y SSO: 100 % de éxito, pero con muy pocos lanzamientos.",
                f"La órbita más frecuente es {top} ({counts.get(top)} lanzamientos) con {pct(rates.get(top), 0)} de éxito."
                if top else "", "GTO, clave para satélites comerciales, es la más difícil."]
    else:
        d.image(s, "eda_success_by_orbit.png", MARGIN, CONTENT_TOP, Inches(8.2), Inches(5.2), viz_hint)
        side = ["Compara la proporción de aterrizajes exitosos por órbita.",
                "Se completa automáticamente con results/eda.json."]
    d.rich(s, Inches(9.2), CONTENT_TOP + Inches(0.3), Inches(3.5), Inches(4.8), [x for x in side if x], size=16,
           space_after=14)

    for title, img, takeaway, notes, kicker in [
        ("Número de vuelo vs. tipo de órbita", "eda_flight_vs_orbit.png",
         ["En LEO, el éxito parece aumentar con el número de vuelo.",
          "En GTO no se aprecia relación entre experiencia y éxito.",
          "Las órbitas nuevas (VLEO) aparecen en los vuelos más recientes."],
         "Combina experiencia y órbita: no todas las órbitas se benefician igual de la experiencia.",
         "EDA Results · Flight Number vs. Orbit Type"),
        ("Carga útil vs. tipo de órbita", "eda_payload_vs_orbit.png",
         ["Con cargas pesadas, los aterrizajes exitosos predominan en Polar, LEO e ISS.",
          "En GTO se mezclan éxitos y fallos en todo el rango de masa.",
          "VLEO concentra cargas pesadas (constelación Starlink)."],
         "Para algunas órbitas, más masa no impide aterrizar; para GTO la relación no es clara.",
         "EDA Results · Payload vs. Orbit Type"),
    ]:
        eda_image_slide(d, title, img, takeaway, viz_hint, notes, kicker)

    s = d.content_slide("Tendencia anual de la tasa de éxito", kicker="EDA Results · Launch Success Yearly Trend", notes=(
        "La línea muestra el porcentaje de aterrizajes exitosos por año. Desde 2013 la tendencia es claramente "
        "ascendente: SpaceX aprendió a aterrizar sus cohetes."),
        source="Fuente: results/eda.json (notebook 05)" if eda else None)
    if eda:
        yearly = eda["success_rate_by_year"]
        d.line_chart(s, list(yearly), list(yearly.values()), MARGIN, CONTENT_TOP, Inches(8.2), Inches(5.2))
        years = list(yearly)
        side = [f"{years[0]}: {pct(yearly[years[0]], 0)} de éxito.",
                f"{years[-1]}: {pct(yearly[years[-1]], 0)} de éxito.",
                "La tasa crece de forma sostenida desde 2013."]
    else:
        d.image(s, "eda_yearly_trend.png", MARGIN, CONTENT_TOP, Inches(8.2), Inches(5.2), viz_hint)
        side = ["Evolución de la tasa de éxito por año.", "Se completa automáticamente con results/eda.json."]
    d.rich(s, Inches(9.2), CONTENT_TOP + Inches(0.3), Inches(3.5), Inches(4.8), side, size=16, space_after=14)

    # --- SQL: una diapositiva por consulta
    sql_insights = [
        "Hay 4 sitios: dos en Cabo Cañaveral (CCAFS LC-40 y SLC-40), uno en Kennedy (KSC LC-39A) y uno en Vandenberg (VAFB SLC-4E).",
        "Los primeros lanzamientos desde Cabo Cañaveral llevaban cápsulas Dragon de prueba y misiones CRS para la NASA.",
        "Masa total transportada para el programa de reabastecimiento de la NASA (CRS).",
        "Masa media que transportaba la versión F9 v1.1 del booster.",
        "El primer aterrizaje exitoso en tierra llegó más de cinco años después del primer lanzamiento.",
        "Boosters de la versión FT que aterrizaron en barcaza con cargas medianas.",
        "Las misiones casi nunca fallan: el reto real es recuperar el booster.",
        "Los boosters Block 5 son los que transportaron la carga máxima.",
        "En 2015 los intentos de aterrizaje en barcaza fallaban: SpaceX aún estaba aprendiendo.",
        "Entre 2010 y 2017 dominan “No attempt” y los intentos en barcaza: era una etapa experimental.",
    ]
    for i, res in enumerate(sql[:10]):
        sql_slide(d, i, res, SQL_QUERIES[i], sql_insights[i], sql_source)

    # ------------------------------------------------------------------ Sección 3: Folium
    section(d, "3", "Análisis de proximidad", "Mapas interactivos con Folium")
    folium_hint = "Se genera con el notebook 06 + scripts/capture_screenshots.py"
    map_slide(d, "Ubicación de los sitios de lanzamiento", "map_launch_sites.png", folium_hint, [
        "Los 4 sitios están en la costa: 3 en Florida y 1 en California.",
        "Están lo más al sur posible de EE. UU. continental: cerca del ecuador la rotación terrestre ayuda a ganar velocidad.",
    ], "Mapa con los cuatro sitios marcados. Señalar que todos están junto al mar.",
       kicker="Launch Sites Proximities Analysis · All Launch Sites on a Map")
    outcomes_text = ["Marcadores verdes = aterrizó; rojos = no aterrizó.", "Los clústeres agrupan los lanzamientos de cada sitio."]
    if folium_res:
        best_site = max(folium_res["site_outcomes"], key=lambda r: r["tasa_exito"])
        outcomes_text.append(f"Mayor tasa de éxito: {best_site['Launch Site']} ({pct(best_site['tasa_exito'], 1)}).")
    map_slide(d, "Resultados de lanzamiento por sitio", "map_launch_outcomes.png", folium_hint, outcomes_text,
              "Cada marcador es un lanzamiento coloreado según su resultado.",
              kicker="Launch Sites Proximities Analysis · Launch Outcomes by Site")
    prox_text = ["Distancias medidas con la fórmula del haversine."]
    if folium_res:
        prox_text = [f"{name}: {km:.2f} km".replace(".", ",") for name, km in folium_res["proximities_km"].items()]
        prox_text.append("Cerca de costa, tren y carretera; lejos de las ciudades por seguridad.")
    map_slide(d, f"Entorno de {folium_res['proximities_site'] if folium_res else 'CCAFS SLC-40'}",
              "map_proximities.png", folium_hint, prox_text,
              "Distancias a costa, ferrocarril, autopista y ciudad. El sitio está a muy poca distancia del mar y de las vías de transporte, y lejos de núcleos urbanos.",
              kicker="Launch Sites Proximities Analysis · Proximities of a Launch Site")

    # ------------------------------------------------------------------ Sección 4: Dashboard
    section(d, "4", "Dashboard con Plotly Dash", "Exploración interactiva de los lanzamientos")
    dash_hint = "Se genera con scripts/capture_screenshots.py"
    pie_text = ["Cuota de lanzamientos exitosos que aporta cada sitio."]
    if dash:
        share = dash["by_site"][dash["top_share_site"]]["cuota_exitos"]
        pie_text = [f"{dash['top_share_site']} aporta la mayor parte de los éxitos ({pct(share)}).",
                    "Cuota de éxitos por sitio sobre el total de aterrizajes exitosos."]
    map_slide(d, "Éxitos por sitio de lanzamiento", "dash_pie_all_sites.png", dash_hint, pie_text,
              "Gráfico de pastel con todos los sitios.",
              kicker="Plotly Dash Dashboard · Launch Success Count for All Sites")
    top_text = ["Sitio con la mayor tasa de éxito: éxito frente a fallo."]
    if dash:
        top = dash["top_rate_site"]
        top_text = [f"{top} es el sitio con mayor tasa de éxito ({pct(dash['by_site'][top]['tasa_exito'])}).",
                    f"{int(dash['by_site'][top]['exitos'])} de {int(dash['by_site'][top]['total'])} lanzamientos aterrizaron."]
    map_slide(d, "Sitio con mayor tasa de éxito", "dash_pie_top_site.png", dash_hint, top_text,
              "Al elegir un sitio en el desplegable, el pastel muestra su proporción de éxito y fallo.",
              kicker="Plotly Dash Dashboard · Launch Site with Highest Success Ratio")

    if dash:
        bins = {k: v for k, v in dash["success_rate_by_payload_bin"].items() if v["rate"] is not None}
        best_bin = max(bins, key=lambda k: bins[k]["rate"])
        boosters = dash["success_rate_by_booster"]
        best_booster = max(boosters, key=lambda k: boosters[k]["rate"])
        scatter_text = [f"Versión de booster con mayor tasa de éxito: {best_booster} ({pct(boosters[best_booster]['rate'], 0)}, "
                        f"{boosters[best_booster]['launches']} lanzamiento(s)).",
                        "Cada color es una versión del booster; arriba aterrizó (1), abajo no (0)."]
        range_text = [f"Rango de carga con mayor tasa de éxito: {best_bin} kg ({pct(bins[best_bin]['rate'], 0)}, "
                      f"{bins[best_bin]['launches']} lanzamientos).",
                      "El control deslizante filtra la masa y el gráfico se actualiza al instante."]
    else:
        scatter_text = ["Relación entre masa de carga y resultado, por versión del booster."]
        range_text = ["El control deslizante filtra la masa y el gráfico se actualiza al instante."]
    map_slide(d, "Carga útil vs. resultado: todos los sitios", "dash_scatter_all.png", dash_hint, scatter_text,
              "Dispersión con todo el rango de masa. Cada color es una versión del booster.",
              kicker="Plotly Dash Dashboard · Payload vs. Launch Outcome (All Sites)")
    map_slide(d, "Carga útil vs. resultado: 2.000–6.000 kg", "dash_scatter_range.png", dash_hint, range_text,
              "La misma dispersión tras reducir el rango de masa con el control deslizante.",
              kicker="Plotly Dash Dashboard · Payload vs. Launch Outcome (Range Slider)")

    # ------------------------------------------------------------------ Sección 5: ML
    section(d, "5", "Análisis predictivo", "Clasificación del aterrizaje con machine learning")
    s = d.content_slide("Exactitud de clasificación", kicker="Predictive Analysis Results · Classification Accuracy", notes=(
        "Gris: exactitud media en validación cruzada (dentro del entrenamiento). Naranja: exactitud sobre los "
        "lanzamientos de prueba que el modelo nunca vio. Con tan pocos datos de prueba, varios modelos empatan."),
        source=model_source)
    names = list(models["models"])
    d.bar_chart(s, names, {
        "Validación cruzada": [models["models"][n]["cv_accuracy"] for n in names],
        "Prueba": [models["models"][n]["test_accuracy"] for n in names],
    }, MARGIN, CONTENT_TOP, Inches(7.4), Inches(5.2), colors=[SOFT, ORANGE], legend=True, y_max=1.1)
    rows = [[n, ", ".join(f"{k}={_fmt(v)}" for k, v in list(models["models"][n]["best_params"].items())[:3])]
            for n in names]
    d.table(s, ["Modelo", "Mejores hiperparámetros"], rows, Inches(8.3), CONTENT_TOP, Inches(4.43),
            col_widths=[1.4, 2.6], size=12, row_h=Inches(0.62))
    d.rich(s, Inches(8.3), CONTENT_TOP + Inches(3.75), Inches(4.43), Inches(1.4), [
        [("Mejor modelo: ", True), (f"{best} ({pct(test_acc, 2)} en prueba).", False)],
        [("Ojo: ", True), (f"con {models['test_size']} casos de prueba, un acierto vale {point} puntos; "
                           f"en validación cruzada los 4 modelos quedan entre {cv_range}.", False)],
    ], size=14)

    s = d.content_slide(f"Matriz de confusión: {best}", kicker="Predictive Analysis Results · Confusion Matrix", notes=(
        "La matriz compara lo que predijo el modelo con lo que pasó de verdad. La diagonal (verde) son aciertos; "
        "fuera de la diagonal están los errores."),
        source=model_source)
    confusion_grid(d, s, best_m["confusion_matrix"], MARGIN + Inches(1.6), CONTENT_TOP + Inches(0.5))
    (tn, fp), (fn, tp) = best_m["confusion_matrix"]
    total = tn + fp + fn + tp
    d.rich(s, Inches(7.6), CONTENT_TOP + Inches(0.4), Inches(5.1), Inches(4.8), [
        [("Aciertos: ", True), (f"{tn + tp} de {total} lanzamientos ({pct((tn + tp) / total, 2)}).", False)],
        [("Verdaderos positivos (VP): ", True), (f"{tp} — aterrizó y el modelo lo predijo.", False)],
        [("Verdaderos negativos (VN): ", True), (f"{tn} — no aterrizó y el modelo lo predijo.", False)],
        [("Falsos positivos (FP): ", True), (f"{fp} — el modelo predijo aterrizaje y no ocurrió.", False)],
        [("Falsos negativos (FN): ", True), (f"{fn} — aterrizó pero el modelo predijo que no.", False)],
        [("Lectura: ", True), (error_reading(fp, fn), False)],
    ], size=16, space_after=14)

    # ------------------------------------------------------------------ Conclusiones
    s = d.new_slide(dark=True, notes=(
        "Cerrar con las respuestas a las preguntas del inicio. El aprendizaje principal: la experiencia, la órbita "
        "y el sitio influyen en el aterrizaje, y un modelo sencillo acierta la gran mayoría de los casos de prueba."))
    d.text(s, "CONCLUSIONS", MARGIN, Inches(0.3), CONTENT_W, Inches(0.3), size=12, bold=True, color=ORANGE)
    d.text(s, "Conclusiones", MARGIN, Inches(0.55), CONTENT_W, Inches(0.75), size=30, bold=True, color=WHITE,
           anchor=MSO_ANCHOR.MIDDLE)
    conclusions = [
        ("Experiencia", "Cuantos más vuelos acumula SpaceX, mayor es la tasa de aterrizaje: el éxito crece año tras año desde 2013."),
        ("Órbita", "Las órbitas con 100 % de éxito tienen pocas misiones; GTO, la más comercial, es la más difícil."),
        ("Sitio", f"{dash['top_rate_site'] if dash else 'KSC LC-39A'} destaca por su tasa de éxito; todos los sitios están en la costa, cerca de infraestructura y lejos de ciudades."),
        ("Predicción", f"El mejor modelo ({best}) alcanza {pct(test_acc, 2)} de exactitud en prueba. {main_error(best_m)}"),
    ]
    for i, (head, body) in enumerate(conclusions):
        col, row = divmod(i, 2)
        x = MARGIN + col * Inches(6.15)
        y = CONTENT_TOP + Inches(0.1) + row * Inches(2.5)
        d.box(s, x, y, Inches(5.9), Inches(2.2), fill=NAVY_2)
        d.badge(s, str(i + 1), x + Inches(0.3), y + Inches(0.3), d=Inches(0.55), size=16)
        d.text(s, head, x + Inches(1.05), y + Inches(0.3), Inches(4.6), Inches(0.55), size=20, bold=True,
               color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, body, x + Inches(1.05), y + Inches(0.95), Inches(4.6), Inches(1.15), size=16, color=WHITE)
    d.footer(s, model_source, dark=True)

    s = d.content_slide("Ideas innovadoras y próximos pasos", kicker="Innovative Insights", notes=(
        "Propuestas para llevar el proyecto más allá del curso. Son ideas, no resultados: hay que presentarlas "
        "como trabajo futuro."))
    ideas = [
        ("Más datos", "Añadir los lanzamientos posteriores a 2020: con solo 18 casos de prueba cada acierto vale 5,6 puntos."),
        ("Variables nuevas", "Meteorología del día del lanzamiento y estado del mar en la zona de la barcaza."),
        ("Probabilidades", "Calibrar el modelo para estimar la probabilidad de aterrizaje y, con ella, el coste esperado."),
        ("Validación temporal", "Entrenar con el pasado y probar con el futuro para medir la capacidad real de predicción."),
        ("Interpretabilidad", "Explicar cada predicción (importancia de variables, SHAP) para los equipos de negocio."),
        ("Producto", "Publicar el modelo como API y conectarlo al dashboard para simular ofertas en licitaciones."),
    ]
    for i, (head, body) in enumerate(ideas):
        col, row = divmod(i, 3)
        x = MARGIN + col * Inches(6.15)
        y = CONTENT_TOP + Inches(0.05) + row * Inches(1.75)
        d.badge(s, str(i + 1), x, y, d=Inches(0.55), size=16)
        d.text(s, head, x + Inches(0.8), y - Inches(0.02), Inches(5.1), Inches(0.45), size=18, bold=True, color=NAVY)
        d.text(s, body, x + Inches(0.8), y + Inches(0.45), Inches(5.1), Inches(1.1), size=16)

    s = d.content_slide("Apéndice: repositorio y reproducibilidad", kicker="Appendix", notes=(
        "Dónde está cada parte del trabajo y cómo reproducirlo."))
    d.table(s, ["Entregable", "Archivo en el repositorio"], [
        ["Recolección con API", NB["api"]], ["Web scraping", NB["scraping"]], ["Data wrangling", NB["wrangling"]],
        ["EDA con SQL", NB["sql"]], ["EDA con visualización", NB["viz"]], ["Mapas con Folium", NB["folium"]],
        ["Dashboard con Plotly Dash", NB["dash"]], ["Machine learning", NB["ml"]],
        ["Generador de esta presentación", "presentation/generate_presentation.py"],
    ], MARGIN, CONTENT_TOP, Inches(7.8), col_widths=[2.4, 4.0], size=14, row_h=Inches(0.48))
    d.text(s, "Reproducir el proyecto", Inches(8.9), CONTENT_TOP, Inches(3.8), Inches(0.45), size=20, bold=True,
           color=NAVY)
    d.rich(s, Inches(8.9), CONTENT_TOP + Inches(0.6), Inches(3.8), Inches(3.6), [
        "pip install -r requirements.txt",
        "Ejecutar los notebooks 01 → 07",
        "python scripts/capture_screenshots.py",
        "python presentation/generate_presentation.py",
    ], size=15, space_after=12)
    d.text(s, REPO_URL, Inches(8.9), Inches(5.6), Inches(3.8), Inches(0.4), size=14, color=ORANGE, url=REPO_URL)

    d.save(OUTPUT)
    return OUTPUT


# ======================================================================================
# Plantillas de diapositiva reutilizables
# ======================================================================================
def section(d: Deck, number: str, title: str, subtitle: str) -> None:
    s = d.new_slide(dark=True, notes=f"Sección {number}: {title}.")
    d.badge(s, number, MARGIN, Inches(2.4), d=Inches(1.2), size=40)
    d.text(s, title, MARGIN, Inches(3.85), Inches(11), Inches(1.0), size=44, bold=True, color=WHITE)
    d.text(s, subtitle, MARGIN, Inches(4.85), Inches(11), Inches(0.6), size=20, color=SOFT)


def flow_slide(d: Deck, title: str, steps: list[str], nb_key: str, notes: str, kicker: str | None = None) -> None:
    """Diagrama de flujo en dos filas de tres pasos."""
    s = d.content_slide(title, notes=notes, kicker=kicker)
    box_w, box_h = Inches(3.5), Inches(1.45)
    gap_x = (CONTENT_W - 3 * box_w) / 2
    for i, step in enumerate(steps):
        row, col = divmod(i, 3)
        if row == 1:
            col = 2 - col  # la segunda fila avanza de derecha a izquierda
        x = MARGIN + col * (box_w + gap_x)
        y = CONTENT_TOP + Inches(0.35) + row * Inches(2.25)
        last = i == len(steps) - 1
        d.box(s, x, y, box_w, box_h, fill=NAVY if last else TINT)
        d.badge(s, str(i + 1), x + Inches(0.25), y + (box_h - Inches(0.55)) // 2, d=Inches(0.55), size=16)
        d.text(s, step.split("\n"), x + Inches(0.95), y, box_w - Inches(1.1), box_h, size=15,
               color=WHITE if last else INK, anchor=MSO_ANCHOR.MIDDLE, bold=last)
        if last:
            continue
        if row == 0 and col < 2:
            d.arrow(s, x + box_w + Inches(0.05), y + box_h // 2, x + box_w + gap_x - Inches(0.05), y + box_h // 2)
        elif row == 0:
            d.arrow(s, x + box_w // 2, y + box_h + Inches(0.05), x + box_w // 2, y + Inches(2.25) - Inches(0.05))
        else:
            d.arrow(s, x - Inches(0.05), y + box_h // 2, x - gap_x + Inches(0.05), y + box_h // 2)
    d.link_chip(s, nb_key, MARGIN, Inches(6.4))


def eda_image_slide(d: Deck, title: str, img: str, takeaway: list[str], hint: str, notes: str, kicker: str) -> None:
    s = d.content_slide(title, notes=notes, kicker=kicker)
    d.image(s, img, MARGIN, CONTENT_TOP, Inches(8.2), Inches(5.2), hint)
    d.rich(s, Inches(9.2), CONTENT_TOP + Inches(0.3), Inches(3.5), Inches(4.8), takeaway, size=16, space_after=14)


def map_slide(d: Deck, title: str, img: str, hint: str, takeaway: list[str], notes: str,
              kicker: str | None = None) -> None:
    s = d.content_slide(title, notes=notes, kicker=kicker)
    d.image(s, img, MARGIN, CONTENT_TOP, Inches(8.4), Inches(5.2), hint)
    d.rich(s, Inches(9.4), CONTENT_TOP + Inches(0.3), Inches(3.33), Inches(4.8), takeaway, size=16, space_after=14)


def sql_slide(d: Deck, i: int, res: dict, query: str, insight: str, source: str) -> None:
    s = d.content_slide(res["task"], notes=(
        f"Consulta SQL {i + 1}. A la izquierda la consulta, a la derecha el resultado. Idea clave: {insight}"),
        source=source, kicker=f"EDA with SQL Results · Query {i + 1}/10 · {SQL_TASKS_EN[i]}")
    code_h = Inches(3.9)
    d.code(s, query, MARGIN, CONTENT_TOP + Inches(0.4), Inches(5.1), code_h)
    columns, rows = res["columns"], res["rows"]
    max_rows = 12
    shown = rows[:max_rows]
    tx = MARGIN + Inches(5.5)
    tw = CONTENT_W - Inches(5.5)
    widths = [max(len(str(c)), *(len(str(r[j])) for r in shown)) if shown else len(str(c))
              for j, c in enumerate(columns)]
    widths = [min(max(w, len(str(c)) + 2, 6), 40) for w, c in zip(widths, columns)]
    n = len(shown)
    row_h = Inches(0.42) if n <= 6 else Inches(0.32)
    size = 14 if n <= 6 and len(columns) <= 3 else 11
    d.table(s, columns, shown, tx, CONTENT_TOP + Inches(0.4), tw, col_widths=widths, size=size, row_h=row_h)
    if len(rows) > max_rows:
        d.text(s, f"… y {len(rows) - max_rows} filas más", tx, Inches(6.1), tw, Inches(0.3), size=11, color=MUTED)
    d.box(s, MARGIN, Inches(5.85), Inches(5.1), Inches(0.95), fill=TINT)
    d.text(s, insight, MARGIN + Inches(0.2), Inches(5.85), Inches(4.7), Inches(0.95), size=14, color=INK,
           anchor=MSO_ANCHOR.MIDDLE)


def confusion_grid(d: Deck, s, cm, x, y, cell=Inches(1.9)) -> None:
    (tn, fp), (fn, tp) = cm
    labels = [[("VN", tn, GREEN), ("FP", fp, RED)], [("FN", fn, RED), ("VP", tp, GREEN)]]
    d.text(s, "Predicción", x, y - Inches(0.5), 2 * cell, Inches(0.4), size=14, bold=True, color=NAVY,
           align=PP_ALIGN.CENTER)
    for j, name in enumerate(["No aterriza", "Aterriza"]):
        d.text(s, name, x + j * cell, y - Inches(0.15), cell, Inches(0.3), size=12, color=MUTED,
               align=PP_ALIGN.CENTER)
    for i, name in enumerate(["No aterrizó", "Aterrizó"]):
        d.text(s, name, x - Inches(1.4), y + Inches(0.2) + i * cell, Inches(1.3), cell - Inches(0.2), size=12,
               color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(2):
        for j in range(2):
            tag, value, color = labels[i][j]
            cx, cy = x + j * cell, y + Inches(0.2) + i * cell
            d.box(s, cx + Inches(0.05), cy + Inches(0.05), cell - Inches(0.1), cell - Inches(0.1),
                        fill=color if i == j else TINT, radius=0.08)
            fg = WHITE if i == j else color
            d.text(s, str(value), cx, cy + Inches(0.35), cell, Inches(0.8), size=44, bold=True, color=fg,
                   align=PP_ALIGN.CENTER)
            d.text(s, tag, cx, cy + Inches(1.2), cell, Inches(0.4), size=14, bold=True, color=fg,
                   align=PP_ALIGN.CENTER)
    label = d.text(s, "Valor real", x - Inches(1.6) - cell, y + Inches(0.2) + cell - Inches(0.2), 2 * cell,
                   Inches(0.4), size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    label.rotation = 270


def error_reading(fp: int, fn: int) -> str:
    if fp == fn == 0:
        return "el modelo acierta todos los casos de prueba."
    if fp >= fn:
        return "reconoce bien los aterrizajes; le cuesta más identificar los lanzamientos que no aterrizan."
    return "identifica bien los fallos; le cuesta más reconocer los aterrizajes."


def main_error(model: dict) -> str:
    (_, fp), (fn, _) = model["confusion_matrix"]
    if fp == fn == 0:
        return "No comete errores en el conjunto de prueba."
    return "Su error más frecuente es el falso positivo." if fp >= fn else "Su error más frecuente es el falso negativo."


def _fmt(value) -> str:
    if isinstance(value, float):
        return f"{value:.4g}"
    return str(value)


def export_pdf(pptx: Path) -> Path | None:
    """Convierte la presentación a PDF con LibreOffice, si está instalado.

    El curso pide entregar el PDF con el nombre "Data Science Capstone Project Report".
    Sin LibreOffice, abre el .pptx en PowerPoint y usa Archivo > Exportar > PDF.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice no está instalado: exporta el PDF desde PowerPoint.")
        return None
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(pptx)],
                       check=True, capture_output=True, timeout=300)
        target = pptx.parent / PDF_NAME
        shutil.move(str(Path(tmp) / f"{pptx.stem}.pdf"), target)
    return target


if __name__ == "__main__":
    out = build()
    print(f"Presentación generada: {out.relative_to(ROOT)}")
    if "--pdf" in sys.argv:
        pdf = export_pdf(out)
        if pdf:
            print(f"PDF generado: {pdf.relative_to(ROOT)}")
